"""Generate travel_log_template.xlsx - the starter travel log to fill in."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).parent / "travel_log_template.xlsx"

HEADERS = ["Country", "Start Date", "End Date"]
EXAMPLE_ROWS = [
    ["France", "2024-06-01", "2024-06-14"],
    ["Italy", "2024-06-15", "2024-06-22"],
]


def build() -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Travel Log"

    header_fill = PatternFill("solid", fgColor="1F3864")
    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(EXAMPLE_ROWS, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    sheet.freeze_panes = "A2"

    note = DataValidation(type="textLength", operator="greaterThan", formula1="0")
    note.prompt = "Dates must be YYYY-MM-DD. Ranges should not overlap."
    note.promptTitle = "Travel log"
    sheet.add_data_validation(note)
    note.add(f"A2:C{len(EXAMPLE_ROWS) + 200}")

    workbook.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"Written: {build()}")
